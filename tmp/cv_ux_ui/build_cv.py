from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
DEST = ROOT / "output/docx/David_Cruz_CV_Diseno_UX_UI.docx"
DEST.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(0.58), Inches(0.58)
section.left_margin = section.right_margin = Inches(0.68)
section.header_distance, section.footer_distance = Inches(0.25), Inches(0.25)
CONTENT_WIDTH = section.page_width - section.left_margin - section.right_margin


def style_font(style, size, bold=False, color="20252B"):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().append(OxmlElement("w:lang"))
    style.element.rPr[-1].set(qn("w:val"), "es-SV")


style_font(doc.styles["Normal"], 11.25)
normal = doc.styles["Normal"].paragraph_format
normal.line_spacing = Pt(15)
normal.space_after = Pt(4)
normal.widow_control = True

style_font(doc.styles["Title"], 24, True, "000000")
doc.styles["Title"].paragraph_format.space_after = Pt(4)
doc.styles["Title"].paragraph_format.line_spacing = Pt(29)
style_font(doc.styles["Subtitle"], 12.5, True, "000000")
doc.styles["Subtitle"].paragraph_format.space_after = Pt(9)
doc.styles["Subtitle"].paragraph_format.line_spacing = Pt(16)
style_font(doc.styles["Heading 1"], 11, True, "000000")
hfmt = doc.styles["Heading 1"].paragraph_format
hfmt.space_before, hfmt.space_after = Pt(15), Pt(7)
hfmt.line_spacing = Pt(14)
hfmt.keep_with_next = True

for name, size, bold in [("Company", 11.5, True), ("Role", 11, False), ("Contact", 10, False), ("Project", 11.5, True)]:
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = doc.styles["Normal"]
    style_font(s, size, bold, "000000" if bold else "20252B")
    s.paragraph_format.keep_with_next = name != "Contact"
    s.paragraph_format.space_after = Pt(2 if name != "Contact" else 1)
    s.paragraph_format.line_spacing = Pt(14 if name != "Contact" else 13)

style_font(doc.styles["List Bullet"], 11.25)
bfmt = doc.styles["List Bullet"].paragraph_format
bfmt.left_indent, bfmt.first_line_indent = Inches(0.14), Inches(-0.14)
bfmt.space_after, bfmt.line_spacing = Pt(4), Pt(15)
bfmt.keep_together = True


def link(paragraph, text, url):
    relationship = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "76263F")
    properties.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    properties.append(size)
    run.append(properties)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def heading(text):
    return doc.add_paragraph(text.upper(), "Heading 1")


def job(company, role, dates, bullets):
    p = doc.add_paragraph(style="Company")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(company)
    r = p.add_run("\t" + dates)
    r.bold, r.font.size = False, Pt(10)
    doc.add_paragraph(role, "Role")
    for text in bullets:
        doc.add_paragraph(text, "List Bullet")


def project(title, description, url):
    p = doc.add_paragraph(style="Project")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(title + "\t")
    link(p, "Ver en Figma", url)
    doc.add_paragraph(description)


doc.core_properties.title = "David Cruz Curriculum Vitae Diseño UX UI"
doc.core_properties.subject = "Candidatura a puestos de Diseñador UX UI"
doc.core_properties.author = "David Alejandro Cruz Cruz"
doc.core_properties.keywords = "UX/UI, Figma, diseño de interfaces, sistemas de diseño, frontend"

doc.add_paragraph("David Alejandro Cruz Cruz", "Title")
doc.add_paragraph("DISEÑADOR UX UI", "Subtitle")
p = doc.add_paragraph(style="Contact")
p.add_run("San Salvador Centro, El Salvador  |  +503 7415-0122  |  ")
link(p, "dcruzer92@gmail.com", "mailto:dcruzer92@gmail.com")
p = doc.add_paragraph(style="Contact")
p.add_run("Portafolio UX/UI: ")
link(p, "dcruzsv.vercel.app", "https://dcruzsv.vercel.app")
p.add_run("  |  LinkedIn: ")
link(p, "linkedin.com/in/cruzsv", "https://www.linkedin.com/in/cruzsv")

heading("Perfil profesional")
doc.add_paragraph(
    "Diseñador UX/UI e Ingeniero en Sistemas y Computación con experiencia en interfaces para plataformas empresariales y productos de software. "
    "Mi experiencia incluye diseño en Figma, rediseño de vistas y definición de un sistema de diseño unificado. "
    "Mi base en desarrollo web me permite convertir decisiones de diseño en componentes funcionales, con atención a la usabilidad, la accesibilidad y la consistencia visual."
)
heading("Competencias UX UI")
doc.add_paragraph(
    "Diseño de interfaces en Figma · Sistemas de diseño · Flujos de interacción · "
    "Componentes reutilizables · Usabilidad y accesibilidad · Visualización de datos"
)
heading("Experiencia relevante")
job(
    "Software Mobile Solutions S.A. de C.V.", "Analista Programador Web", "may 2025 - actualidad",
    [
        "Rediseñé las vistas del producto principal y definí un sistema de diseño UX/UI unificado, estandarizando colores, botones, componentes visuales y patrones de interacción.",
        "Reestructuré la experiencia e interfaz de distintos módulos con criterios de diseño centrado en el usuario, accesibilidad y rendimiento.",
        "Implementé interfaces interactivas con componentes reutilizables, validaciones en tiempo real y gráficos mediante Kendo UI, HighCharts y jQuery.",
        "Diseñé y desarrollé un sistema de tickets con tablero Kanban, asignación de responsables, evidencias y estados controlados para dar seguimiento a las incidencias.",
        "Diseñé y desarrollé un asistente conversacional de soporte con respuestas en gráficas, tablas y mapas, además de análisis de archivos adjuntos.",
    ],
)
job(
    "CodeEye Solutions Inc.", "Desarrollador de Software Junior", "abr 2023 - jun 2024",
    [
        "Diseñé y desarrollé interfaces de usuario utilizando Figma como herramienta principal de diseño.",
        "Propuse e implementé mejoras de interfaz y experiencia de usuario con Bootstrap y JavaScript.",
        "Colaboré en IRIS Code Scanner, una extensión de detección de vulnerabilidades para Visual Studio Marketplace integrada al producto principal de la empresa.",
    ],
)

doc.add_page_break()
heading("Proyectos de portafolio")
project(
    "Panel de Vulnerabilidades",
    "Interfaz en Figma para organizar espacios de trabajo, vulnerabilidades y reportes. Diseño de indicadores, tablas y gráficos para la consulta de información.",
    "https://www.figma.com/design/KEbpsM6D6wlBzbGeSDJd5l/Vulnerability-Management-Dashboard?node-id=0-1",
)
project(
    "Aplicación Neuroamigable",
    "Concepto de aplicación con pantallas de registro, evaluación, recursos y perfil. Interfaz orientada a una navegación clara y a la comodidad del usuario.",
    "https://www.figma.com/design/g039NhW9ZAy8NcF2knXqZk/MindMap",
)

heading("Experiencia adicional")
job(
    "Universidad Tecnológica de El Salvador", "Docente Hora Clase", "jul 2026 - sep 2026",
    ["Impartí Pensamiento Lógico y Algoritmos I a estudiantes de Diseño Gráfico, Informática e Ingeniería."],
)
job(
    "Tecnologías 101 S.A. de C.V.", "Desarrollador de Software Junior", "oct 2024 - nov 2024",
    ["Mejoré interfaces y nuevas funcionalidades durante la migración de un sistema en Laravel; apoyé su localización e internacionalización."],
)
job(
    "Caja Mutual de los Empleados del Ministerio de Educación", "Asistente de Soporte Técnico Computacional", "feb 2019 - ago 2019",
    ["Realicé mantenimiento preventivo de equipos informáticos y trabajé en el cableado de redes LAN."],
)

heading("Formación académica")
doc.add_paragraph("Universidad Tecnológica de El Salvador", "Company")
for degree, dates in [
    ("Preespecialidad en Ingeniería de Inteligencia Artificial", "mar - dic 2025"),
    ("Ingeniería en Sistemas y Computación", "2020 - 2024"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(degree)
    r = p.add_run("\t" + dates)
    r.font.size = Pt(10)

heading("Herramientas y colaboración")
for label, value in [
    ("Diseño", "Figma, sistemas de diseño y componentes de interfaz."),
    ("Frontend", "HTML, CSS, JavaScript, Bootstrap, jQuery, Kendo UI y HighCharts."),
    ("Colaboración", "Kanban, GitHub, GitLab y Bitbucket."),
    ("Inglés", "lectura técnica, nivel básico A1."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(label + ": ").bold = True
    p.add_run(value)

style_font(doc.styles["Footer"], 9, color="555555")
doc.styles["Footer"].paragraph_format.tab_stops.clear_all()
footer = section.footer.paragraphs[0]
footer.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
footer.paragraph_format.space_after = Pt(0)
footer.add_run("David Cruz · Diseño UX UI\t")
for kind in ["PAGE", "NUMPAGES"]:
    if kind == "NUMPAGES":
        footer.add_run(" / ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), kind)
    footer._p.append(field)
for r in footer.runs:
    r.font.name, r.font.size, r.font.color.rgb = "Calibri", Pt(9), RGBColor.from_string("555555")

for xml_root in (doc.element, doc.styles.element):
    for border in xml_root.xpath(".//w:pBdr"):
        border.getparent().remove(border)

doc.save(DEST)
print(DEST)
